#!/usr/bin/env python3
"""
REAL-TIME MONITORING DASHBOARD
==============================
View scraper progress in real-time.

Usage:
    python monitor.py

This reads leads.xlsx and progress.json to show accurate stats.
"""

import json
import time
import os
from datetime import datetime
from pathlib import Path

# Files to monitor
PROGRESS_FILE = 'progress.json'
STATS_FILE = 'monitoring_stats.json'
EXCEL_INPUT = 'leads.xlsx'


def load_json(file_path):
    """Load JSON file safely"""
    try:
        if Path(file_path).exists():
            with open(file_path, 'r') as f:
                return json.load(f)
    except Exception as e:
        return {}
    return {}


def count_total_addresses():
    """Count total addresses in leads.xlsx"""
    try:
        import openpyxl
        if not Path(EXCEL_INPUT).exists():
            print(f"⚠️ Warning: {EXCEL_INPUT} not found!")
            return 0
            
        wb = openpyxl.load_workbook(EXCEL_INPUT, read_only=True, data_only=True)
        ws = wb.active
        
        # Count non-empty rows in address column (usually column B)
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            # Check if there's an address (usually in second column)
            if row and len(row) > 1 and row[1]:  # Column B
                count += 1
        
        wb.close()
        return count
    except Exception as e:
        print(f"⚠️ Error reading {EXCEL_INPUT}: {e}")
        return 0


def count_phones_found():
    """Count phone numbers in leads.xlsx Column C (non-N/A values)"""
    phone_count = 0
    try:
        import openpyxl
        if not Path(EXCEL_INPUT).exists():
            return 0
            
        wb = openpyxl.load_workbook(EXCEL_INPUT, read_only=True, data_only=True)
        ws = wb.active
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 3:
                continue
            
            # Column C (index 2) - count non-empty, non-N/A values
            phone = row[2]
            if phone and str(phone).strip().upper() != 'N/A':
                phone_count += 1
        
        wb.close()
    except Exception as e:
        pass
    
    return phone_count


def clear_screen():
    """Clear terminal screen"""
    os.system('cls' if os.name == 'nt' else 'clear')


def print_dashboard():
    """Print the monitoring dashboard"""
    clear_screen()
    
    # Load data
    progress = load_json(PROGRESS_FILE)
    stats = load_json(STATS_FILE)
    
    completed = len(progress.get('completed', []))
    total = count_total_addresses()
    phone_count = count_phones_found()
    
    # Calculate stats
    if total > 0:
        percent = (completed / total) * 100
        remaining = total - completed
        success_rate = (phone_count / completed * 100) if completed > 0 else 0
    else:
        percent = 0
        remaining = 0
        success_rate = 0
    
    # Estimate completion
    if completed > 0 and 'started_at' in stats:
        try:
            start = datetime.fromisoformat(stats['started_at'])
            elapsed = (datetime.now() - start).total_seconds() / 3600  # hours
            rate = completed / elapsed if elapsed > 0 else 0
            eta_hours = remaining / rate if rate > 0 else 0
            eta_days = eta_hours / 24
        except:
            rate = 0
            eta_days = 0
    else:
        rate = 0
        eta_days = 0
    
    # Print dashboard
    print("=" * 80)
    print("📊 TRUEPEOPLESEARCH SCRAPER - REAL-TIME DASHBOARD")
    print("=" * 80)
    print(f"🕐 Current Time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print()
    
    # Progress
    print("📈 PROGRESS:")
    print(f"   Total in leads.xlsx: {total:,} addresses")
    print(f"   Attempted: {completed:,} addresses")
    print(f"   Phones Found: {phone_count:,} ({success_rate:.1f}% success)")
    print(f"   Remaining: {remaining:,}")
    print(f"   Progress: {percent:.1f}%")
    
    # Progress bar
    bar_width = 50
    filled = int(bar_width * percent / 100) if percent > 0 else 0
    bar = '█' * filled + '░' * (bar_width - filled)
    print(f"   [{bar}] {percent:.1f}%")
    print()
    
    # Performance
    print("⚡ PERFORMANCE:")
    print(f"   Processing Rate: {rate:.1f} addresses/hour")
    print(f"   Est. Daily Rate: {rate * 24:.0f} addresses/day")
    if eta_days > 0:
        print(f"   Est. Time Remaining: {eta_days:.1f} days")
    print()
    
    # System stats
    if stats:
        print("🖥️ SYSTEM:")
        print(f"   Total Runs: {stats.get('total_runs', 0)}")
        print(f"   Total Crashes: {stats.get('total_crashes', 0)}")
        if stats.get('last_crash'):
            print(f"   Last Crash: {stats['last_crash'][:19]}")
        print()
    
    print("=" * 80)
    print("Results are saved to: leads.xlsx (Column C)")
    print("Press Ctrl+C to exit | Auto-refresh every 10 seconds")
    print("=" * 80)


def main():
    """Main monitoring loop"""
    print("Starting monitoring dashboard...\n")
    time.sleep(1)
    
    try:
        while True:
            print_dashboard()
            time.sleep(10)  # Refresh every 10 seconds
    except KeyboardInterrupt:
        clear_screen()
        print("\n👋 Monitoring stopped. Goodbye!\n")


if __name__ == '__main__':
    main()
